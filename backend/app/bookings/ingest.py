"""Ingest historical bookings into the canonical `taxi_bookings` table from CSV / XLSX / JSON / API / database.

Column mapping is configurable (config/bookings.yaml → `column_map`) so any dispatch system export can be loaded
without code changes. Idempotent on booking_id."""
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Iterator

from dateutil import parser as dtparser
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import load_config
from app.db.models import TaxiBooking
from app.processors.normalize import normalize_postcode

CANONICAL = ("booking_id", "booking_created_at", "pickup_datetime", "dropoff_datetime", "pickup_postcode", "dropoff_postcode",
             "pickup_lat", "pickup_lon", "dropoff_lat", "dropoff_lon", "fare", "status", "vehicle_type")


def column_map() -> dict[str, str]:
    """canonical_field → source column name. Defaults to identity mapping."""
    cfg = load_config("bookings.yaml") or {}
    m = {k: k for k in CANONICAL}
    m.update({k: v for k, v in (cfg.get("column_map") or {}).items() if v})
    return m


def _dt(v: Any) -> datetime | None:
    if v in (None, "", "NULL"):
        return None
    if isinstance(v, datetime):
        return v
    try:
        return dtparser.parse(str(v), dayfirst=True)
    except (ValueError, OverflowError):
        return None


def _f(v: Any) -> float | None:
    try:
        return float(v) if v not in (None, "") else None
    except (TypeError, ValueError):
        return None


def to_canonical(row: dict[str, Any], cmap: dict[str, str] | None = None) -> dict[str, Any] | None:
    cmap = cmap or column_map()
    g = lambda k: row.get(cmap.get(k, k))
    pickup = _dt(g("pickup_datetime"))
    bid = g("booking_id")
    if not bid or not pickup:
        return None
    return {
        "booking_id": str(bid)[:80], "booking_created_at": _dt(g("booking_created_at")), "pickup_datetime": pickup,
        "dropoff_datetime": _dt(g("dropoff_datetime")), "pickup_postcode": normalize_postcode(g("pickup_postcode")),
        "dropoff_postcode": normalize_postcode(g("dropoff_postcode")), "pickup_lat": _f(g("pickup_lat")), "pickup_lon": _f(g("pickup_lon")),
        "dropoff_lat": _f(g("dropoff_lat")), "dropoff_lon": _f(g("dropoff_lon")), "fare": _f(g("fare")),
        "status": (str(g("status") or "")[:20] or None), "vehicle_type": (str(g("vehicle_type") or "")[:30] or None),
    }


# ---- readers ---------------------------------------------------------------------------------------------------------
def read_csv(path: Path) -> Iterator[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as fh:
        yield from csv.DictReader(fh)


def read_json(path: Path) -> Iterator[dict[str, Any]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    rows = data if isinstance(data, list) else data.get("bookings") or data.get("data") or []
    for r in rows:
        if isinstance(r, dict):
            yield r


def read_xlsx(path: Path) -> Iterator[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # optional dependency
        raise RuntimeError("openpyxl is required for XLSX ingest: pip install openpyxl") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, [])]
    for r in rows:
        yield {header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}


def read_api(url: str, headers: dict[str, str] | None = None) -> Iterator[dict[str, Any]]:
    """Generic JSON API connector (GET, optional bearer). Response must be a list or {bookings|data: [...]}. SSRF-validated."""
    import httpx
    from app.core.netsafe import validate_url
    validate_url(url, None)
    r = httpx.get(url, headers=headers or {}, timeout=60)
    r.raise_for_status()
    data = r.json()
    rows = data if isinstance(data, list) else data.get("bookings") or data.get("data") or []
    for row in rows:
        if isinstance(row, dict):
            yield row


def read_database(url: str, query: str) -> Iterator[dict[str, Any]]:
    """Database connector: any SQLAlchemy URL + SELECT returning the (mapped) canonical columns."""
    from sqlalchemy import create_engine, text
    eng = create_engine(url)
    with eng.connect() as conn:
        for row in conn.execute(text(query)).mappings():
            yield dict(row)


def read_path(path: str | Path) -> Iterator[dict[str, Any]]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".csv":
        return read_csv(p)
    if ext == ".json":
        return read_json(p)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(p)
    raise ValueError(f"unsupported booking file type: {ext}")


# ---- loader ----------------------------------------------------------------------------------------------------------
def ingest(db: Session, rows: Iterable[dict[str, Any]], source_label: str | None = None) -> dict[str, int]:
    cmap = column_map()
    existing = set(db.scalars(select(TaxiBooking.booking_id)).all())
    stats = {"read": 0, "created": 0, "updated": 0, "rejected": 0}
    for raw in rows:
        stats["read"] += 1
        rec = to_canonical(raw, cmap)
        if not rec:
            stats["rejected"] += 1
            continue
        if rec["booking_id"] in existing:
            obj = db.scalar(select(TaxiBooking).where(TaxiBooking.booking_id == rec["booking_id"]))
            for k, v in rec.items():
                setattr(obj, k, v)
            stats["updated"] += 1
        else:
            db.add(TaxiBooking(**rec, source_file=source_label))
            existing.add(rec["booking_id"])
            stats["created"] += 1
    db.commit()
    return stats
